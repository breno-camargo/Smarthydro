import os
import calendar
import logging
from datetime import datetime
from dateutil.relativedelta import relativedelta
import pyodbc
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

from core.database import build_connection_string
from core.config_manager import load_config, get_base_dir

logger = logging.getLogger(__name__)

MESES_ABREV = [
    "Jan", "Fev", "Mar", "Abr", "Mai", "Jun",
    "Jul", "Ago", "Set", "Out", "Nov", "Dez"
]

MESES_COMPLETO = [
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"
]


def get_cycle_dates(target_year, target_month):
    """
    Retorna as datas de início e fim do ciclo de faturamento para o mês/ano alvo.
    Fim: 28 do target_month às 23:59:59.
    Início: 29 do mês anterior às 00:00:00 (ajustando para 28 no caso de fevereiro se necessário).
    """
    d_fim = datetime(target_year, target_month, 28, 23, 59, 59)
    prev_dt = datetime(target_year, target_month, 1) - relativedelta(days=1)
    py, pm = prev_dt.year, prev_dt.month
    day_ini = min(29, calendar.monthrange(py, pm)[1])
    d_ini = datetime(py, pm, day_ini, 0, 0, 0)
    return d_ini, d_fim


def fetch_annual_history(config=None, target_date=None, progress_callback=None):
    """
    Consulta o banco de dados SQL Server e retorna os dados de consumo dos últimos 12 ciclos fechados.
    Retorna um dicionário com lista 'months' e sumário 'kpis'.
    """
    if config is None:
        config = load_config()

    if target_date is None:
        target_date = datetime.now()

    # Se hoje ainda for antes do dia 28 do mês corrente, o último ciclo fechado é o do mês anterior
    if target_date.day < 28:
        last_closed_month_dt = datetime(target_date.year, target_date.month, 1) - relativedelta(months=1)
    else:
        last_closed_month_dt = datetime(target_date.year, target_date.month, 1)

    valor_m3 = float(config.get("default_m3_price", 63.68))
    conn_str, driver = build_connection_string(config)

    # Monta a lista dos 12 ciclos em ordem cronológica (do mais antigo para o mais recente)
    cycles_to_query = []
    for i in range(11, -1, -1):
        m_dt = last_closed_month_dt - relativedelta(months=i)
        d_ini, d_fim = get_cycle_dates(m_dt.year, m_dt.month)
        cycles_to_query.append((m_dt.year, m_dt.month, d_ini, d_fim))

    conn = pyodbc.connect(conn_str, timeout=15)
    cur = conn.cursor()

    sql_query = """
    SET NOCOUNT ON;
    SET DATEFORMAT ymd;
    SELECT 
        SUM(ISNULL(CASE 
            WHEN (l_fim.val - l_ini.val) < 0 THEN 0.0
            ELSE (l_fim.val - l_ini.val) / 1000.0
        END, 0.0)) AS Total_m3,
        COUNT(*) as Total_Medidores
    FROM tbTrendLogRelation r
    CROSS APPLY (
        SELECT TOP(1) FloatVALUE as val 
        FROM tbLogTimeValues 
        WHERE ParentID = r.EntityID AND DateTimeStamp BETWEEN ? AND ? 
        ORDER BY DateTimeStamp ASC
    ) l_ini
    CROSS APPLY (
        SELECT TOP(1) FloatVALUE as val 
        FROM tbLogTimeValues 
        WHERE ParentID = r.EntityID AND DateTimeStamp BETWEEN ? AND ? 
        ORDER BY DateTimeStamp DESC
    ) l_fim
    WHERE r.Type = 'trend.ETLog' AND r.Name NOT LIKE '%Sala 1001-1006%';
    """

    months_data = []
    prev_consumo = None

    for idx, (yr, mo, ini, fim) in enumerate(cycles_to_query):
        if progress_callback:
            try:
                progress_callback(idx + 1, len(cycles_to_query), f"Consultando {MESES_ABREV[mo-1]}/{str(yr)[-2:]}...")
            except Exception:
                pass

        ini_str = ini.strftime("%Y-%m-%d %H:%M:%S")
        fim_str = fim.strftime("%Y-%m-%d %H:%M:%S")

        cur.execute(sql_query, (ini_str, fim_str, ini_str, fim_str))
        row = cur.fetchone()

        consumo_m3 = round(float(row[0] or 0.0), 1)
        total_medidores = int(row[1] or 0)
        valor_rs = round(consumo_m3 * valor_m3, 2)

        # Cálculo de variação vs mês anterior
        if prev_consumo is not None and prev_consumo > 0:
            diff_m3 = round(consumo_m3 - prev_consumo, 1)
            diff_pct = round(((consumo_m3 - prev_consumo) / prev_consumo) * 100.0, 1)
        else:
            diff_m3 = 0.0
            diff_pct = 0.0

        prev_consumo = consumo_m3

        mes_label = f"{MESES_ABREV[mo-1]}/{str(yr)[-2:]}"
        mes_extenso = f"{MESES_COMPLETO[mo-1]} de {yr}"
        periodo_label = f"{ini.strftime('%d/%m/%Y')} a {fim.strftime('%d/%m/%Y')}"

        months_data.append({
            "ano": yr,
            "mes_num": mo,
            "mes_label": mes_label,
            "mes_extenso": mes_extenso,
            "periodo": periodo_label,
            "consumo_m3": consumo_m3,
            "valor_rs": valor_rs,
            "medidores": total_medidores,
            "diff_m3": diff_m3,
            "diff_pct": diff_pct
        })

    conn.close()

    # Cálculo dos KPIs consolidados
    total_consumo = sum(m["consumo_m3"] for m in months_data)
    total_faturamento = sum(m["valor_rs"] for m in months_data)
    qtd_meses = len(months_data) or 1
    media_mensal_m3 = round(total_consumo / qtd_meses, 1)
    media_mensal_rs = round(total_faturamento / qtd_meses, 2)

    # Identificar pico e mínimo
    pico = max(months_data, key=lambda x: x["consumo_m3"]) if months_data else None
    minimo = min(months_data, key=lambda x: x["consumo_m3"]) if months_data else None

    kpis = {
        "total_consumo_m3": round(total_consumo, 1),
        "total_faturamento_rs": round(total_faturamento, 2),
        "media_mensal_m3": media_mensal_m3,
        "media_mensal_rs": media_mensal_rs,
        "pico_mes": pico["mes_label"] if pico else "-",
        "pico_m3": pico["consumo_m3"] if pico else 0.0,
        "pico_periodo": pico["periodo"] if pico else "-",
        "minimo_mes": minimo["mes_label"] if minimo else "-",
        "minimo_m3": minimo["consumo_m3"] if minimo else 0.0,
        "tarifa_m3": valor_m3,
        "periodo_geral": f"{months_data[0]['periodo'].split(' a ')[0]} a {months_data[-1]['periodo'].split(' a ')[1]}" if months_data else ""
    }

    return {
        "months": months_data,
        "kpis": kpis
    }


def export_annual_history_excel(history_data, output_path):
    """
    Gera um relatório anual executivo em Excel (.xlsx) com tabela formatada e gráfico de barras CompaSSS.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Histórico Anual"
    ws.views.sheetView[0].showGridLines = True

    # Cores e Estilos
    COLOR_DARK_GREEN = "3D6B24"
    COLOR_LIGHT_GREEN = "EBF3E6"
    COLOR_ACCENT = "90C671"
    COLOR_BORDER = "D5E5C9"

    font_title = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    font_sub = Font(name="Segoe UI", size=9, italic=True, color="EBF3E6")
    font_kpi_label = Font(name="Segoe UI", size=8, bold=True, color="55664C")
    font_kpi_val = Font(name="Segoe UI", size=13, bold=True, color="3D6B24")
    font_th = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
    font_td = Font(name="Segoe UI", size=9, color="1B2A12")
    font_td_bold = Font(name="Segoe UI", size=9, bold=True, color="1B2A12")

    fill_header = PatternFill(start_color=COLOR_DARK_GREEN, end_color=COLOR_DARK_GREEN, fill_type="solid")
    fill_kpi = PatternFill(start_color="FAFCF8", end_color="FAFCF8", fill_type="solid")
    fill_alt = PatternFill(start_color=COLOR_LIGHT_GREEN, end_color=COLOR_LIGHT_GREEN, fill_type="solid")
    fill_total = PatternFill(start_color="D5E5C9", end_color="D5E5C9", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=COLOR_BORDER),
        right=Side(style='thin', color=COLOR_BORDER),
        top=Side(style='thin', color=COLOR_BORDER),
        bottom=Side(style='thin', color=COLOR_BORDER)
    )

    # 1. Cabeçalho Institucional
    ws.merge_cells("A1:F1")
    cell_title = ws["A1"]
    cell_title.value = "SmartHydro — Histórico Anual de Consumo de Água"
    cell_title.font = font_title
    cell_title.fill = fill_header
    cell_title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:F2")
    cell_sub = ws["A2"]
    kpis = history_data.get("kpis", {})
    cell_sub.value = f"Condomínio Praça Pamplona • Período: {kpis.get('periodo_geral', '')} • Tarifa Base: R$ {kpis.get('tarifa_m3', 63.68):.2f}/m³"
    cell_sub.font = font_sub
    cell_sub.fill = fill_header
    cell_sub.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    # 2. Cards de KPIs (Linhas 4 a 5)
    kpi_defs = [
        ("A4:A5", "Consumo Anual Total", f"{kpis.get('total_consumo_m3', 0):,.1f} m³".replace(",", "X").replace(".", ",").replace("X", ".")),
        ("B4:B5", "Faturamento Total", f"R$ {kpis.get('total_faturamento_rs', 0):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")),
        ("C4:C5", "Média Mensal", f"{kpis.get('media_mensal_m3', 0):,.1f} m³".replace(",", "X").replace(".", ",").replace("X", ".")),
        ("D4:D5", "Mês de Maior Consumo", f"{kpis.get('pico_mes', '-')} ({kpis.get('pico_m3', 0):,.1f} m³)".replace(",", "X").replace(".", ",").replace("X", ".")),
        ("E4:F5", "Mês Mais Econômico", f"{kpis.get('minimo_mes', '-')} ({kpis.get('minimo_m3', 0):,.1f} m³)".replace(",", "X").replace(".", ",").replace("X", "."))
    ]

    for coord, label, val in kpi_defs:
        top_cell = coord.split(":")[0]
        ws[top_cell] = f"{label}\n{val}"
        ws[top_cell].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws[top_cell].fill = fill_kpi
        ws[top_cell].font = font_kpi_val
        ws[top_cell].border = thin_border
        if ":" in coord:
            ws.merge_cells(coord)

    ws.row_dimensions[4].height = 22
    ws.row_dimensions[5].height = 22

    # 3. Tabela de Dados Mensais (Linha 7)
    headers = [
        ("Mês / Ano", 14),
        ("Período de Medição", 24),
        ("Consumo (m³)", 16),
        ("Variação (m³)", 14),
        ("Variação (%)", 14),
        ("Valor Faturado (R$)", 20)
    ]

    th_row = 7
    ws.row_dimensions[th_row].height = 22
    for c_idx, (h_text, width) in enumerate(headers, 1):
        cell = ws.cell(row=th_row, column=c_idx, value=h_text)
        cell.font = font_th
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center" if c_idx > 2 else "left", vertical="center")
        cell.border = thin_border
        col_letter = get_column_letter(c_idx)
        ws.column_dimensions[col_letter].width = width

    months = history_data.get("months", [])
    current_row = 8

    for idx, m in enumerate(months):
        ws.row_dimensions[current_row].height = 20
        row_fill = fill_alt if idx % 2 == 1 else PatternFill(fill_type=None)

        # Col 1: Mês
        c1 = ws.cell(row=current_row, column=1, value=m["mes_label"])
        c1.font = font_td_bold
        c1.alignment = Alignment(horizontal="left", vertical="center")
        c1.border = thin_border
        if row_fill.fill_type:
            c1.fill = row_fill

        # Col 2: Período
        c2 = ws.cell(row=current_row, column=2, value=m["periodo"])
        c2.font = font_td
        c2.alignment = Alignment(horizontal="center", vertical="center")
        c2.border = thin_border
        if row_fill.fill_type:
            c2.fill = row_fill

        # Col 3: Consumo m³
        c3 = ws.cell(row=current_row, column=3, value=m["consumo_m3"])
        c3.font = font_td_bold
        c3.number_format = "#,##0.0"
        c3.alignment = Alignment(horizontal="right", vertical="center")
        c3.border = thin_border
        if row_fill.fill_type:
            c3.fill = row_fill

        # Col 4: Variação m³
        v_m3 = m["diff_m3"]
        c4 = ws.cell(row=current_row, column=4, value=v_m3 if idx > 0 else "-")
        c4.font = font_td
        if idx > 0:
            c4.number_format = "+#,##0.0;-#,##0.0;0.0"
        c4.alignment = Alignment(horizontal="right", vertical="center")
        c4.border = thin_border
        if row_fill.fill_type:
            c4.fill = row_fill

        # Col 5: Variação %
        v_pct = m["diff_pct"]
        txt_pct = f"{v_pct:+.1f}%" if idx > 0 else "-"
        c5 = ws.cell(row=current_row, column=5, value=txt_pct)
        c5.font = font_td
        c5.alignment = Alignment(horizontal="right", vertical="center")
        c5.border = thin_border
        if row_fill.fill_type:
            c5.fill = row_fill

        # Col 6: Valor R$
        c6 = ws.cell(row=current_row, column=6, value=m["valor_rs"])
        c6.font = font_td
        c6.number_format = "R$ #,##0.00"
        c6.alignment = Alignment(horizontal="right", vertical="center")
        c6.border = thin_border
        if row_fill.fill_type:
            c6.fill = row_fill

        current_row += 1

    # Linha Total Acumulada
    ws.row_dimensions[current_row].height = 22
    t1 = ws.cell(row=current_row, column=1, value="TOTAL ANUAL")
    t1.font = font_td_bold
    t1.fill = fill_total
    t1.border = thin_border

    t2 = ws.cell(row=current_row, column=2, value="12 Meses Consolidados")
    t2.font = font_td_bold
    t2.fill = fill_total
    t2.alignment = Alignment(horizontal="center", vertical="center")
    t2.border = thin_border

    t3 = ws.cell(row=current_row, column=3, value=f"=SUM(C8:C{current_row-1})")
    t3.font = font_td_bold
    t3.number_format = "#,##0.0"
    t3.fill = fill_total
    t3.alignment = Alignment(horizontal="right", vertical="center")
    t3.border = thin_border

    t4 = ws.cell(row=current_row, column=4, value="")
    t4.fill = fill_total
    t4.border = thin_border

    t5 = ws.cell(row=current_row, column=5, value="")
    t5.fill = fill_total
    t5.border = thin_border

    t6 = ws.cell(row=current_row, column=6, value=f"=SUM(F8:F{current_row-1})")
    t6.font = font_td_bold
    t6.number_format = "R$ #,##0.00"
    t6.fill = fill_total
    t6.alignment = Alignment(horizontal="right", vertical="center")
    t6.border = thin_border

    # 4. Gráfico de Barras Embutido
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Evolução do Consumo Mensal de Água (m³)"
    chart.y_axis.title = "Consumo (m³)"
    chart.x_axis.title = "Mês de Medição"
    chart.legend = None

    data_ref = Reference(ws, min_col=3, min_row=7, max_row=current_row - 1)
    cats_ref = Reference(ws, min_col=1, min_row=8, max_row=current_row - 1)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)

    chart.width = 18
    chart.height = 12
    ws.add_chart(chart, f"A{current_row + 2}")

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)

    try:
        from core.report_generator import _inject_cached_formula_values
        tot_c = round(sum(float(m.get("consumo_m3") or 0) for m in months), 1)
        tot_v = round(sum(float(m.get("valor_rs") or 0) for m in months), 2)
        _inject_cached_formula_values(output_path, {
            f"C{current_row}": tot_c,
            f"F{current_row}": tot_v,
        })
    except Exception as e:
        logging.warning(f"Não foi possível injetar cache no histórico anual: {e}")

    return output_path
